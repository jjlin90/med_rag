# -*- coding: utf-8 -*-
"""
同步 artifacts/presentation.pptx 中与代码不一致的表述。
该演示文件是本地产物，不随仓库分发；新克隆请用 --pptx 指定自己的文件。

背景：PPT 生成于双通道改造之前，以下 12 处已与代码脱节：
  - 检索/重排参数（Top8 / Top4 → Top-16 / Top-5 子块 / Top-2 父块）
  - 混合检索三向量（实际未启用 ColBERT）
  - 路由架构（意图预判 + 医疗保护 + 关键词守卫 → 双通道 + softmax 归一化）
  - 分层降级（新增 L0/L1/L2）

设计：只替换**完整匹配**的 run.text，不触碰段落与字体级格式；
      执行后打印命中计数，未命中项会明确报 MISS，避免静默漏改。

用法：
  python scripts/update_pptx_text.py                 # 默认改 artifacts/presentation.pptx
  python scripts/update_pptx_text.py --dry-run       # 只预览，不写盘
  python scripts/update_pptx_text.py --pptx path/to/slides.pptx --dry-run
"""

import argparse
import shutil
import sys
from pathlib import Path

BASE_DIR = Path(__file__).parent.parent
DEFAULT_PPTX = BASE_DIR / "artifacts" / "presentation.pptx"

# 旧文本 -> 新文本（必须与 run.text 完全一致才会被替换）
REPLACEMENTS = {
    # ---- P10 在线检索流程六步 ----
    "Milvus 混合检索，粗排 Top8 候选。":
        "Milvus 混合检索（过滤下推只召回子块），粗排 Top-16。",
    "BGE-reranker 交叉编码精排 Top4。":
        "BGE-reranker 交叉编码精排 Top-2 父块。",

    # ---- P11 策略与重排 ----
    "BGE-reranker 精排：CrossEncoder 对 Top8 候选精细打分，输出 Top4。":
        "BGE-reranker 精排：CrossEncoder 对候选父块精细打分，输出 Top-2。",
    "混合检索：dense（语义）+ sparse（关键词）+ multi-vector 三者融合。":
        "混合检索：dense（语义 1.0）+ sparse（词权 0.7）加权融合，ColBERT 未启用。",

    # ---- P13 智能路由（架构已重构）----
    "用户提问 → BERT 意图分类 → medical 走 RAG；general 走 FAQ / 直答。":
        "用户提问 → FAQ 快通道优先（Redis → BM25 归一化 ≥0.85）→ 未命中降级 RAG 深通道。",
    "bert-base-chinese 微调，输出 general / medical 及置信度。":
        "bert-base-chinese 微调，输出 general / medical；只在深通道内执行一次。",
    "医疗保护":
        "双通道架构",
    "医疗问题跳过 FAQ 直走 RAG，保证权威出处。":
        "所有问题先过 FAQ 快通道，未命中才降级 RAG 深通道。",
    "FAQ 守卫":
        "归一化阈值",
    "查询须与 FAQ 共享有效关键词，否则回退 RAG。":
        "BM25 分 softmax 归一到 (0,1]，≥ 0.85 才算命中。",

    # ---- P16 演示页：实际精排输出 Top-2 ----
    "📎 检索来源（Top 3）":
        "📎 检索来源（Top 2）",

    # ---- P17 工程成果 ----
    "修复 FAQ 误答：医疗问题跳过 FAQ + 关键词守卫。":
        "修复 FAQ 误答：BM25 softmax 归一化（阈值 0.85）；检索失败时 L0/L1/L2 分层降级，不让 LLM 凭空作答。",
}


def main():
    parser = argparse.ArgumentParser(description="同步 PPT 中与代码脱节的表述")
    parser.add_argument("--pptx", default=str(DEFAULT_PPTX), help="目标 pptx 路径；默认文件仅在本地存在，不随仓库分发")
    parser.add_argument("--dry-run", action="store_true", help="只预览不写盘")
    args = parser.parse_args()

    pptx_path = Path(args.pptx)
    if not pptx_path.is_file():
        print(f"[FAIL] 文件不存在: {pptx_path}")
        print("PPT 是本地产物，不随仓库分发。请用 --pptx 指定已有 .pptx 文件。")
        return 1

    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    hit = {k: 0 for k in REPLACEMENTS}

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text in REPLACEMENTS:
                        hit[run.text] += 1
                        run.text = REPLACEMENTS[run.text]

    # 校验：逐项报告命中情况，未命中要显式暴露
    print("=" * 60)
    ok = miss = 0
    for old, new in REPLACEMENTS.items():
        if hit[old]:
            print(f"  [OK x{hit[old]}] {old[:34]}... -> {new[:34]}...")
            ok += 1
        else:
            print(f"  [MISS]    {old}")
            miss += 1
    print("=" * 60)

    if args.dry_run:
        print(f"[DRY-RUN] 未写盘。命中 {ok} 项 / 缺失 {miss} 项")
        return 0

    if ok:
        backup = pptx_path.with_suffix(".bak.pptx")
        if not backup.exists():
            shutil.copy2(pptx_path, backup)
            print(f"[备份] {backup.name}")
        prs.save(str(pptx_path))
        print(f"[完成] 已写入 {pptx_path.name}：替换 {ok} 项 / 缺失 {miss} 项")
    else:
        print("[跳过] 无任何命中，未写盘")

    return 0 if miss == 0 else 2


if __name__ == "__main__":
    sys.exit(main())
