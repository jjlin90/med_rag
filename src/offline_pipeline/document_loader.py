"""
Document Loader Module
支持多格式文档加载（docx/pdf/pptx/md/txt）；OCR 接口尚为占位
"""

import os
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation
from pathlib import Path
from typing import List, Dict, Any
import logging

from ..config.settings import Config

logger = logging.getLogger(__name__)

# 文件加载映射
LOADER_MAPPING = {
    '.pdf': 'pdf',
    '.docx': 'docx',
    '.pptx': 'pptx',
    '.txt': 'txt',
    '.md': 'markdown'
}

class Document:
    """统一文档对象"""
    def __init__(self, page_content: str, metadata: Dict[str, Any]):
        self.page_content = page_content
        self.metadata = metadata

    def __repr__(self):
        return f"Document(source={self.metadata.get('source', 'unknown')}, length={len(self.page_content)})"

class DocumentLoader:
    """文档加载器"""

    def __init__(self, config: Config):
        self.config = config
        self.supported_extensions = set(LOADER_MAPPING.keys())

    def load_documents_from_directory(self, directory: Path) -> List[Document]:
        """递归加载目录下所有支持的文档"""
        documents = []

        if not directory.exists():
            logger.warning(f"Directory not found: {directory}")
            return documents

        # 支持递归遍历所有子目录
        for root, dirs, files in os.walk(directory):
            for file in files:
                file_path = Path(root) / file
                ext = file_path.suffix.lower()

                if ext in self.supported_extensions:
                    try:
                        docs = self.load_document(file_path)
                        documents.extend(docs)
                        logger.info(f"Loaded {len(docs)} documents from {file_path}")
                    except Exception as e:
                        logger.error(f"Failed to load {file_path}: {str(e)}")
                        continue

        return documents

    def load_document(self, file_path: Path) -> List[Document]:
        """加载单个文档"""
        ext = file_path.suffix.lower()

        if ext not in LOADER_MAPPING:
            logger.warning(f"Unsupported file type: {ext}")
            return []

        loader_type = LOADER_MAPPING[ext]

        try:
            if loader_type == 'pdf':
                return self._load_pdf(file_path)
            elif loader_type == 'docx':
                return self._load_docx(file_path)
            elif loader_type == 'pptx':
                return self._load_pptx(file_path)
            elif loader_type == 'txt':
                return self._load_txt(file_path)
            elif loader_type == 'markdown':
                return self._load_markdown(file_path)
        except Exception as e:
            logger.error(f"Error loading {file_path}: {str(e)}")
            return []

    def _load_pdf(self, file_path: Path) -> List[Document]:
        """加载PDF文档"""
        documents = []

        try:
            # 使用PyMuPDF加载PDF
            with fitz.open(file_path) as doc:
                for page_num in range(len(doc)):
                    page = doc[page_num]

                    # 提取文本
                    text = page.get_text()

                    # 如果文本为空，尝试OCR
                    if not text.strip():
                        text = self._ocr_page(page)

                    # 元数据
                    metadata = {
                        'source': self._extract_source(file_path),
                        'file_path': str(file_path),
                        'page_number': page_num + 1,
                        'total_pages': len(doc),
                        'timestamp': int(file_path.stat().st_mtime),
                        'file_type': 'pdf'
                    }

                    documents.append(Document(text, metadata))

        except Exception as e:
            logger.error(f"PDF loading error: {str(e)}")
            # 尝试OCR作为备选方案
            try:
                documents = self._load_pdf_with_ocr(file_path)
            except Exception as e2:
                logger.error(f"OCR fallback also failed: {str(e2)}")

        return documents

    def _load_docx(self, file_path: Path) -> List[Document]:
        """加载Word文档"""
        documents = []

        try:
            doc = DocxDocument(file_path)
            full_text = []

            for paragraph in doc.paragraphs:
                full_text.append(paragraph.text)

            text = '\n'.join(full_text)

            metadata = {
                'source': self._extract_source(file_path),
                'file_path': str(file_path),
                'timestamp': int(file_path.stat().st_mtime),
                'file_type': 'docx',
                'paragraph_count': len(doc.paragraphs)
            }

            documents.append(Document(text, metadata))

        except Exception as e:
            logger.error(f"DOCX loading error: {str(e)}")

        return documents

    def _load_pptx(self, file_path: Path) -> List[Document]:
        """加载PowerPoint文档"""
        documents = []

        try:
            prs = Presentation(file_path)
            full_text = []

            for slide in prs.slides:
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text)

                if slide_text:
                    full_text.append('\n'.join(slide_text))

            text = '\n'.join(full_text)

            metadata = {
                'source': self._extract_source(file_path),
                'file_path': str(file_path),
                'timestamp': int(file_path.stat().st_mtime),
                'file_type': 'pptx',
                'slide_count': len(prs.slides)
            }

            documents.append(Document(text, metadata))

        except Exception as e:
            logger.error(f"PPTX loading error: {str(e)}")
            # 尝试OCR作为备选方案
            try:
                documents = self._load_pptx_with_ocr(file_path)
            except Exception as e2:
                logger.error(f"OCR fallback also failed: {str(e2)}")

        return documents

    def _load_txt(self, file_path: Path) -> List[Document]:
        """加载文本文件"""
        documents = []

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()

            metadata = {
                'source': self._extract_source(file_path),
                'file_path': str(file_path),
                'timestamp': int(file_path.stat().st_mtime),
                'file_type': 'txt',
                'char_count': len(text)
            }

            documents.append(Document(text, metadata))

        except UnicodeDecodeError:
            # 尝试其他编码
            try:
                with open(file_path, 'r', encoding='gbk') as f:
                    text = f.read()

                metadata = {
                    'source': self._extract_source(file_path),
                    'file_path': str(file_path),
                    'timestamp': int(file_path.stat().st_mtime),
                    'file_type': 'txt',
                    'char_count': len(text)
                }

                documents.append(Document(text, metadata))

            except Exception as e:
                logger.error(f"TXT loading error: {str(e)}")
        except Exception as e:
            logger.error(f"TXT loading error: {str(e)}")

        return documents

    def _load_markdown(self, file_path: Path) -> List[Document]:
        """加载Markdown文件"""
        documents = []

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()

            metadata = {
                'source': self._extract_source(file_path),
                'file_path': str(file_path),
                'timestamp': int(file_path.stat().st_mtime),
                'file_type': 'markdown',
                'char_count': len(text)
            }

            documents.append(Document(text, metadata))

        except Exception as e:
            logger.error(f"Markdown loading error: {str(e)}")

        return documents

    def _ocr_page(self, page) -> str:
        """OCR处理PDF页面（简化版本，实际使用时需要集成OCR引擎）"""
        # 这里简化处理，实际项目中应该集成如RapidOCR、Tesseract等OCR引擎
        # 返回空字符串表示OCR失败
        return ""

    def _load_pdf_with_ocr(self, file_path: Path) -> List[Document]:
        """使用OCR加载PDF"""
        # 这是一个占位符方法，实际实现需要集成OCR引擎
        logger.warning("OCR not implemented yet for PDF")
        return []

    def _load_pptx_with_ocr(self, file_path: Path) -> List[Document]:
        """使用OCR加载PPTX"""
        # 这是一个占位符方法，实际实现需要集成OCR引擎
        logger.warning("OCR not implemented yet for PPTX")
        return []

    def _extract_source(self, file_path: Path) -> str:
        """从文件路径提取学科来源"""
        # 获取相对于MSD解压目录的路径
        try:
            relative_path = file_path.relative_to(self.config.RAW_MSD_DIR)
            # 取第一个目录名作为学科来源
            parts = relative_path.parts
            if len(parts) > 0:
                return parts[0]
        except ValueError:
            # 如果不在MSD目录下，使用文件名的前缀
            return file_path.stem.split('_')[0] if '_' in file_path.stem else file_path.stem

        return 'unknown'
