"""Boundary regressions using real modules, mocked external services; no model loading."""
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

import numpy as np
from docx import Document as WordDocument
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pydantic import ValidationError

from src.offline_pipeline.document_loader import Document, DocumentLoader
from src.offline_pipeline.data_cleaner import DataCleaner
from src.offline_pipeline.chunk_splitter import Chunk, ChunkSplitter
from src.offline_pipeline.embedding_provider import BGEEmbeddingProvider
from src.offline_pipeline.milvus_store import MilvusStore, MilvusException
from src.online_service.cache_manager import query_cache_key
from src.online_service.faq_search import BM25Index, FAQSearch
from src.online_service.main_api import RAGWebAPI, ChatRequest, QueryRequest, EvaluateItem
from src.online_service.rag_system import RAGSystem
from src.online_service.retrieval import Retrieval, RetrievalResult
from src.online_service.strategy_selector import StrategySelector
from src.online_service.rag_evaluator import RAGEvaluator
from src.online_service.reranker import Reranker


class RegressionTests(unittest.TestCase):
    def test_merged_eval_keeps_zero_rejects_missing(self):
        from scripts.merge_eval_chunks import valid_score, METRICS
        self.assertTrue(valid_score(dict.fromkeys(METRICS, 0.0)))
        for value in (None, float('nan'), float('inf'), -0.1, 1.1, True):
            self.assertFalse(valid_score(dict.fromkeys(METRICS, value)))
        self.assertFalse(valid_score({}))

    def test_merge_eval_report_contract(self):
        from contextlib import redirect_stdout
        import io
        from scripts.merge_eval_chunks import main, METRICS
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            chunk = root / 'chunk.json'
            zero = dict(question='zero', **dict.fromkeys(METRICS, 0.0))
            good = dict(question='good', **dict.fromkeys(METRICS, 1.0))
            chunk.write_text(json.dumps({'engine': 'ragas', 'scores':
                                         [zero, good, zero, {'question': 'missing'}]}), encoding='utf-8')
            output = root / 'out.json'
            argv = ['merge', '--chunk', str(chunk), '--out-json', str(output),
                    '--out-csv', str(root / 'nested/out.csv')]
            with patch('sys.argv', argv), redirect_stdout(io.StringIO()):
                self.assertEqual(main(), 0)
            report = json.loads(output.read_text(encoding='utf-8'))
            self.assertEqual(report['n_total'], 2)
            self.assertEqual(report['n_invalid_dropped'], 1)
            self.assertEqual(report['n_duplicate_dropped'], 1)
            self.assertEqual(report['average']['faithfulness'], 0.5)
            chunk.write_text(json.dumps({'engine': 'fallback', 'scores': [good]}), encoding='utf-8')
            with patch('sys.argv', argv), self.assertRaises(ValueError):
                main()

    def test_cache_preserves_semantics(self):
        for left, right in [('1.5mg', '15mg'), ('120/80', '12080'),
                            ('-5', '5'), ('1 2', '12'), ('mL', 'ML')]:
            with self.subTest(left=left):
                self.assertNotEqual(query_cache_key(left), query_cache_key(right))
        self.assertEqual(query_cache_key(' test '), query_cache_key('test'))

    def test_cache_varies_with_context(self):
        keys = [query_cache_key('q'), query_cache_key('q', 'a'),
                query_cache_key('q', strategy='hyde'),
                query_cache_key('q', history=[{'role': 'user', 'content': 'a'}]),
                query_cache_key('q', history=[{'role': 'user', 'content': 'b'}])]
        self.assertEqual(len(set(keys)), len(keys))

    def test_cleaner_preserves_numbers(self):
        cleaner = DataCleaner(None)
        for value in ['数值120/80，应保留本行', '120/80', '1.5', '5', '第3页说明剂量']:
            self.assertEqual(cleaner._clean_text(value), value)
        self.assertEqual(cleaner._clean_text('第 3 页\nPage 2\n---'), '')

    def test_real_docx_roundtrip(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'test.docx'
            document = WordDocument()
            document.add_paragraph('测试正文 1.5mg')
            document.save(path)
            docs = DocumentLoader(SimpleNamespace(RAW_MSD_DIR=Path(directory))).load_document(path)
            self.assertEqual(len(docs), 1)
            self.assertIn('1.5mg', docs[0].page_content)

    def test_unrelated_single_faq_never_matches(self):
        index = BM25Index()
        index.build_index(['apple'])
        self.assertEqual(index.search_normalized('banana'), [])
        self.assertEqual(index.search_normalized('!!!'), [])
        self.assertGreater(index.search_normalized('apple')[0][0], 0)

    def test_faq_cache_disable_and_offline_cache(self):
        faq = FAQSearch.__new__(FAQSearch)
        faq.cache = Mock()
        faq.cache.get.return_value = {'type': 'faq', 'answer': 'cached'}
        faq.connection = faq.cursor = faq.bm25_index = None
        self.assertEqual(faq.search_faq('q', use_cache=False), (None, True))
        faq.cache.get.assert_not_called()
        faq.cache.set.assert_not_called()
        self.assertEqual(faq.search_faq('q'), ('cached', False))

    def test_strategy_noise_is_deterministic(self):
        selector = StrategySelector(None, None)
        for value in ['', '!!!', 'h', '123']:
            self.assertEqual(selector._normalize_strategy(value), 'direct')
        self.assertEqual(selector._normalize_strategy('"hyde"'), 'hyde')

    def test_source_filter_escapes_input(self):
        source = 'a" or source != "b\\c'
        expression = Retrieval.__new__(Retrieval)._build_source_filter(source)
        self.assertEqual(json.loads(expression.removeprefix('source == ')), source)

    def test_fallback_chunks_cover_text_and_have_unique_ids(self):
        splitter = ChunkSplitter(SimpleNamespace(PARENT_CHUNK_SIZE=20, CHILD_CHUNK_SIZE=8, CHUNK_OVERLAP=2))
        content = 'ab。' + ''.join(chr(0x4e00 + i) for i in range(80))
        parents = splitter._simple_split(Document(content, {'file_path': 'f'}), 20, 'parent')
        children = []
        for parent in parents:
            children.extend(splitter._simple_split(parent, 8, 'child'))
        self.assertEqual(len({c.id for c in children}), len(children))
        self.assertTrue(all(char in ''.join(c.content for c in children) for char in content))
        self.assertTrue(all(len(c.content) <= 8 for c in children))

    def test_bad_overlap_rejected(self):
        with self.assertRaises(ValueError):
            ChunkSplitter(SimpleNamespace(PARENT_CHUNK_SIZE=20, CHILD_CHUNK_SIZE=8, CHUNK_OVERLAP=8))

    def test_collection_count_fallback_is_not_truncated(self):
        from scripts.check_chunk_type_filter import count_by_expr
        collection = Mock()
        collection.query.side_effect = RuntimeError('unsupported count')
        iterator = collection.query_iterator.return_value
        iterator.next.side_effect = [[{'id': 'x'}] * 1000] * 17 + [[]]
        self.assertEqual(count_by_expr(collection, 'parent_id != ""'), 17000)
        iterator.close.assert_called_once()

    def test_split_failure_stops_ingestion(self):
        splitter = ChunkSplitter(SimpleNamespace(PARENT_CHUNK_SIZE=20, CHILD_CHUNK_SIZE=8, CHUNK_OVERLAP=2))
        splitter._split_into_parent_chunks = Mock(side_effect=RuntimeError('invalid document'))
        with self.assertRaisesRegex(RuntimeError, 'invalid document'):
            splitter.split_documents([Document('text', {'file_path': 'bad.md'})])

    def test_ingestion_validates_actual_selected_directory(self):
        from scripts import run_offline_ingest as ingestion
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'src/models/bge-m3').mkdir(parents=True)
            selected = root / 'selected'
            selected.mkdir()
            config = SimpleNamespace(CLEAN_MD_DIR=root / 'missing-default')
            with patch.object(ingestion, 'project_root', root):
                self.assertTrue(ingestion.validate_config(config, selected))
                self.assertFalse(ingestion.validate_config(config))

    def test_cli_ingestion_validation_failure_has_nonzero_exit(self):
        import main
        with patch('scripts.run_offline_ingest.setup_directories'), \
             patch('scripts.run_offline_ingest.validate_config', return_value=False), \
             patch('scripts.run_offline_ingest.process_documents') as process:
            with self.assertRaises(SystemExit) as error:
                main.process_data_mode(Mock(), Path('missing'))
            self.assertEqual(error.exception.code, 1)
            process.assert_not_called()

    def test_legacy_chunk_type_repaired_in_memory(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'chunks.json'
            path.write_text(json.dumps([{'id': 'p', 'content': 'text', 'metadata': {},
                                        'parent_id': None, 'chunk_type': 'child'}]))
            splitter = ChunkSplitter(SimpleNamespace(PARENT_CHUNK_SIZE=20, CHILD_CHUNK_SIZE=8, CHUNK_OVERLAP=2))
            chunk = splitter.load_chunks(path)[0]
            self.assertEqual(chunk.chunk_type, 'parent')
            self.assertEqual(chunk.metadata['chunk_type'], 'parent')

    def test_embedding_failure_propagates(self):
        provider = BGEEmbeddingProvider.__new__(BGEEmbeddingProvider)
        provider.dense_model = Mock()
        provider.dense_model.encode.side_effect = RuntimeError('offline')
        with self.assertRaisesRegex(RuntimeError, 'offline'):
            provider.generate_embeddings(['text'])
        dense, sparse = provider.generate_embeddings([])
        self.assertEqual(len(dense), 0)
        self.assertEqual(sparse, [])

    def test_embedding_count_mismatch_fails_batch(self):
        provider = BGEEmbeddingProvider.__new__(BGEEmbeddingProvider)
        provider.device = 'cpu'
        provider.generate_embeddings = Mock(return_value=(np.array([]), []))
        with self.assertRaises(ValueError):
            provider.batch_process([Chunk('c', 'text', {})])

    def test_sparse_vectors_use_integer_keys(self):
        provider = BGEEmbeddingProvider.__new__(BGEEmbeddingProvider)
        provider.dense_model = Mock()
        provider.dense_model.encode.return_value = {
            'dense_vecs': np.array([[1., 2.]]), 'lexical_weights': [{'12': np.float32(.5)}]}
        _, sparse = provider.generate_embeddings(['text'])
        self.assertEqual(sparse, [{12: .5}])

    def test_sdk_error_reaches_retrieval_l2(self):
        store = MilvusStore.__new__(MilvusStore)
        store.collection_name, store.alias, store.nprobe = 'c', 'a', 16
        store._has_chunk_type = True
        store.client = Mock()
        store.client.hybrid_search.side_effect = MilvusException(message='offline')
        retrieval = Retrieval.__new__(Retrieval)
        retrieval.milvus_store = store
        retrieval.embedding_provider = Mock()
        retrieval.embedding_provider.generate_embeddings.return_value = (np.array([[1., 2.]]), [{1: .5}])
        retrieval.top_k_retrieve, retrieval.top_k_children = 16, 5
        retrieval.sparse_weight, retrieval.dense_weight = .7, 1.
        with patch('src.offline_pipeline.milvus_store.Collection'):
            result = retrieval.search_child_to_parent('q')
        self.assertEqual(result.degrade_level, 2)
        self.assertIn('offline', result.error)
        self.assertIn("parent_id != ''", store.child_filter_expr())

    def make_core(self, result):
        core = RAGSystem.__new__(RAGSystem)
        core.config = SimpleNamespace(TOP_K_RERANK=2, ALLOW_LLM_WHEN_NO_CONTEXT=False)
        core.intent_classifier = Mock()
        core.intent_classifier.predict.return_value = {'intent': 'medical', 'confidence': .9}
        core._retrieve_and_merge = Mock(return_value=result)
        core.reranker = Mock()
        core.reranker.rerank.return_value = result.documents
        core.llm_generator = Mock()
        core.llm_generator.generate_with_context.return_value = 'answer'
        return core

    def test_partial_failure_refuses_without_llm(self):
        merged = RAGSystem._merge_retrieval_results([
            RetrievalResult(documents=[{'id': 'p', 'content': 'evidence'}]),
            RetrievalResult(degraded=True, degrade_level=2, error='offline')])
        core = self.make_core(merged)
        result = core.generate('q', strategy='subquery')
        self.assertEqual(result['degrade_level'], 2)
        core.llm_generator.generate_with_context.assert_not_called()

    def test_empty_rerank_refuses_without_llm(self):
        core = self.make_core(RetrievalResult(documents=[{'id': 'p', 'content': 'evidence'}]))
        core.reranker.rerank.return_value = []
        self.assertEqual(core.generate('q', strategy='direct')['degrade_level'], 2)
        core.llm_generator.generate_with_context.assert_not_called()

    def test_general_llm_failure_is_degraded(self):
        core = self.make_core(RetrievalResult())
        core.intent_classifier.predict.return_value = {'intent': 'general'}
        core.llm_generator.generate_with_context.return_value = None
        self.assertEqual(core.generate('q')['degrade_level'], 2)

    def test_classifier_failure_stops_generation(self):
        for prediction in [{'error': 'offline'}, {}, {'intent': 'unknown'}]:
            core = self.make_core(RetrievalResult())
            core.intent_classifier.predict.return_value = prediction
            result = core.generate('q', strategy='direct')
            self.assertEqual(result['degrade_reason'], 'intent_classifier_unavailable')
            self.assertEqual(result['degrade_level'], 2)
            core._retrieve_and_merge.assert_not_called()
            core.llm_generator.generate_with_context.assert_not_called()

    def make_api(self):
        api = RAGWebAPI.__new__(RAGWebAPI)
        api.cache = Mock()
        api.cache.is_connected.return_value = True
        api.cache.get.return_value = None
        api.faq_search = Mock()
        api.faq_search.search_faq.return_value = (None, True)
        api.conversation_store = Mock()
        api.rag_core = Mock()
        api.rag_core.generate.return_value = {
            'answer': 'answer', 'sources': [], 'confidence': .9,
            'intent': 'medical', 'strategy': 'direct'}
        return api

    def test_cache_hit_persists_history(self):
        api = self.make_api()
        api.cache.get.return_value = {'answer': 'cached'}
        response = api._handle_query('q', None, True, session_id='session')
        self.assertTrue(response.used_cache)
        api.conversation_store.update_session_history.assert_called_once_with('session', 'q', 'cached')

    def test_contextual_requests_skip_faq(self):
        for kwargs in [{'history': [{'role': 'user', 'content': 'context'}]},
                       {'source_filter': 'a'}, {'strategy': 'hyde'}]:
            api = self.make_api()
            api._handle_query('q', **{'source_filter': None, 'use_cache': True, **kwargs})
            api.faq_search.search_faq.assert_not_called()

    def test_cache_disabled_and_l2_not_cached(self):
        api = self.make_api()
        api._handle_query('q', None, use_cache=False)
        api.cache.get.assert_not_called()
        api.cache.set.assert_not_called()
        api.faq_search.search_faq.assert_called_once_with('q', use_cache=False)
        api = self.make_api()
        api.rag_core.generate.return_value.update(degraded=True, degrade_level=2)
        api._handle_query('q', None, True)
        api.cache.set.assert_not_called()

    def test_restored_chat_without_timestamp_via_http(self):
        api = self.make_api()
        api.app = FastAPI()
        api._register_routes()
        with TestClient(api.app) as client:
            response = client.post('/chat', json={'messages': [
                {'role': 'user', 'content': 'previous'},
                {'role': 'assistant', 'content': 'answer'},
                {'role': 'user', 'content': 'followup'}]})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(client.post('/query', json={'question': '  '}).status_code, 422)
            self.assertEqual(client.post('/chat', json={'messages': []}).status_code, 400)

    def test_reranker_error_is_visible(self):
        core = self.make_core(RetrievalResult(documents=[{'id': 'p', 'content': 'evidence'}]))
        core.reranker.rerank.side_effect = RuntimeError('offline')
        self.assertEqual(core.generate('q', strategy='direct')['degrade_reason'], 'reranker_unavailable')
        core.llm_generator.generate_with_context.assert_not_called()

    def test_reranker_invalid_scores_rejected(self):
        reranker = Reranker.__new__(Reranker)
        reranker.model = Mock()
        for scores in [[], [float('nan')]]:
            reranker.model.compute_score.return_value = scores
            with self.assertRaises(ValueError):
                reranker.rerank('q', [{'content': 'text'}])

    def test_failed_judge_scores_not_zero(self):
        for raw in [None, 'invalid', '{"faithfulness": "NaN"}', '{"faithfulness": 2}', '{"faithfulness": true}']:
            self.assertIsNone(RAGEvaluator._parse(raw)['faithfulness'])
        self.assertEqual(RAGEvaluator._parse('{"faithfulness": 0}')['faithfulness'], 0)
        average = RAGEvaluator._aggregate([{'faithfulness': None}, {'faithfulness': .8}])
        self.assertEqual(average['faithfulness'], .8)
        self.assertIsNone(average['answer_relevancy'])

    def test_ragas_scores_validate_range_and_requested_columns(self):
        import pandas as pd
        evaluator = RAGEvaluator.__new__(RAGEvaluator)
        evaluator._build_llm = Mock()
        evaluator._build_embeddings = Mock()
        items = evaluator._normalize([{'question': 'q', 'answer': 'a',
                                      'contexts': ['text'], 'ground_truth': 'truth'}])
        for score in [True, -0.1, 1.1, float('inf'), float('nan'), 0.0]:
            result = Mock()
            result.to_pandas.return_value = pd.DataFrame({'faithfulness': [score],
                                                         'answer_relevancy': [.5]})
            with patch('ragas.evaluate', return_value=result):
                report = evaluator._evaluate_ragas(items)
            self.assertEqual(len(report['metrics']), 4)
            self.assertEqual(report['complete_count'], 0)
            self.assertIsNone(report['average']['context_recall'])
            self.assertEqual(report['valid_counts']['context_precision'], 0)
            if score == 0.0 and not isinstance(score, bool):
                self.assertEqual(report['average']['faithfulness'], 0.0)
            else:
                self.assertIsNone(report['average']['faithfulness'])

    def test_ragas_missing_all_columns_keeps_missing_scores(self):
        import pandas as pd
        evaluator = RAGEvaluator.__new__(RAGEvaluator)
        evaluator._build_llm = Mock()
        evaluator._build_embeddings = Mock()
        result = Mock()
        result.to_pandas.return_value = pd.DataFrame({'user_input': ['q']})
        with patch('ragas.evaluate', return_value=result):
            report = evaluator._evaluate_ragas(evaluator._normalize([{'question': 'q', 'answer': 'a'}]))
        self.assertEqual(report['engine'], 'ragas')
        self.assertEqual(report['metrics'], ['faithfulness', 'answer_relevancy'])
        self.assertEqual(report['valid_counts'], dict.fromkeys(report['metrics'], 0))
        self.assertEqual(report['complete_count'], 0)

    def test_evaluate_accepts_reference_alias(self):
        self.assertEqual(EvaluateItem(question='q', answer='a', reference_answer='truth').ground_truth, 'truth')

    def test_debug_start_does_not_pass_app_object_to_reloader(self):
        api = self.make_api()
        api.app = FastAPI()
        with patch('uvicorn.run') as run:
            api.run(debug=True)
        self.assertFalse(run.call_args.kwargs['reload'])
        self.assertEqual(run.call_args.kwargs['log_level'], 'debug')

    def test_untrained_classifier_never_routes_requests(self):
        from src.online_service.intent_classifier import IntentClassifier
        classifier = IntentClassifier.__new__(IntentClassifier)
        with tempfile.TemporaryDirectory() as directory:
            classifier.model_dir = Path(directory)
            with patch('src.online_service.intent_classifier.BertTokenizer.from_pretrained'), \
                 patch('src.online_service.intent_classifier.BertForSequenceClassification.from_pretrained') as model:
                with self.assertRaisesRegex(RuntimeError, 'Trained intent'):
                    classifier._init_model()
                model.assert_not_called()

    def test_all_nan_metric_column_preserved(self):
        import pandas as pd
        result = Mock()
        result.to_pandas.return_value = pd.DataFrame({'faithfulness': [float('nan')],
                                                      'answer_relevancy': [.5], 'user_input': ['q']})
        keys, _ = RAGEvaluator._extract(result)
        self.assertEqual(keys, ['faithfulness', 'answer_relevancy'])

    def test_parallel_evaluation_accepts_zero_rejects_partial(self):
        from scripts.run_parallel_eval import is_valid, METRICS
        from scripts.run_chunked_eval_v2 import is_valid as valid_v2
        zero = {k: 0.0 for k in METRICS}
        for valid in (is_valid, valid_v2):
            self.assertTrue(valid(zero))
            self.assertFalse(valid({**zero, 'context_recall': None}))
            self.assertFalse(valid({**zero, 'context_recall': float('nan')}))
            self.assertFalse(valid({**zero, 'context_recall': True}))

    def test_composite_rounds_once_from_original_scores(self):
        from scripts.run_parallel_eval import merge_results, METRICS
        from scripts.run_chunked_eval_v2 import merge_results as merge_v2
        scores = dict(zip(METRICS, [.8163338095238095, .5006771428571428,
                                  .8404761904761905, .7619047619047619]))
        for merge in (merge_results, merge_v2):
            self.assertEqual(merge({'q': scores}, [])['weighted_composite'], .7298)

    def test_evaluation_merge_missing_column_roundtrips_without_zero_fill(self):
        import csv
        import importlib
        import io
        from contextlib import redirect_stdout
        for name in ('scripts.run_parallel_eval', 'scripts.run_chunked_eval_v2'):
            module = importlib.import_module(name)
            cases = [
                ({'q': dict.fromkeys(module.METRICS, 0.0)}, 0.0, 0.0),
                ({'q': dict.fromkeys(module.METRICS, .5),
                  'partial': {'faithfulness': None, 'answer_relevancy': .5,
                              'context_precision': .5, 'context_recall': .5}}, .5, .5),
                ({'q': {'faithfulness': .5, 'answer_relevancy': .5,
                        'context_precision': .5}}, None, None),
                ({'q': dict.fromkeys(module.METRICS, None)}, None, None),
            ]
            for collected, recall, composite in cases:
                with self.subTest(module=name, collected=collected), tempfile.TemporaryDirectory() as directory:
                    root = Path(directory)
                    report = module.merge_results(collected, [])
                    self.assertEqual(report['average']['context_recall'], recall)
                    self.assertEqual(report['weighted_composite'], composite)
                    with patch.object(module, 'FINAL_JSON', root / 'report.json'), \
                         patch.object(module, 'FINAL_CSV', root / 'report.csv'), \
                         redirect_stdout(io.StringIO()) as output:
                        module.save_final(report)
                        module.print_summary(report)
                    self.assertEqual(json.loads((root / 'report.json').read_text(encoding='utf-8'))['weighted_composite'], composite)
                    with (root / 'report.csv').open(encoding='utf-8-sig', newline='') as source:
                        average = list(csv.DictReader(source))[-1]
                    self.assertEqual(average['context_recall'], '' if recall is None else str(recall))
                    if composite is None:
                        self.assertIn('缺失', output.getvalue())

    def test_evaluation_merge_rejects_invalid_scores_and_displays_empty_result(self):
        import importlib
        import io
        from contextlib import redirect_stdout
        for name in ('scripts.run_parallel_eval', 'scripts.run_chunked_eval_v2'):
            module = importlib.import_module(name)
            for invalid in (True, -0.1, 1.1, float('nan'), float('inf')):
                with self.subTest(module=name, invalid=invalid):
                    report = module.merge_results({'q': {**dict.fromkeys(module.METRICS, .5),
                                                        'faithfulness': invalid}}, [])
                    self.assertIsNone(report['average']['faithfulness'])
                    self.assertIsNone(report['weighted_composite'])
            with tempfile.TemporaryDirectory() as directory, \
                 patch.object(module, 'FINAL_JSON', Path(directory) / 'empty.json'), \
                 patch.object(module, 'FINAL_CSV', Path(directory) / 'empty.csv'), \
                 redirect_stdout(io.StringIO()) as output:
                report = module.merge_results({}, [])
                module.save_final(report)
                module.print_summary(report)
                self.assertIn(report['error'], output.getvalue())
                self.assertNotIn('0.0000', output.getvalue())

    def test_legacy_chunk_merge_retains_missing_scores(self):
        from scripts.run_chunked_eval import merge_results
        report = merge_results([{'engine': 'ragas', 'metrics': ['faithfulness', 'answer_relevancy'],
                                 'scores': [{'question': 'q', 'faithfulness': 0.0, 'answer_relevancy': True}]}])
        self.assertEqual(report['average']['faithfulness'], 0.0)
        self.assertIsNone(report['average']['answer_relevancy'])
        self.assertEqual(report['valid_counts']['answer_relevancy'], 0)
        self.assertEqual(report['complete_count'], 0)

    def test_retry_merge_preserves_missing_zero_and_engine_contract(self):
        from scripts.merge_eval_retries import main, valid
        from contextlib import redirect_stdout
        import io
        from scripts.run_parallel_eval import METRICS
        for score in [True, float('nan'), -0.1, 1.1]:
            self.assertFalse(valid(score))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            base = {'engine': 'ragas', 'metrics': METRICS,
                    'scores': [{'question': 'q', **dict.fromkeys(METRICS, None)}],
                    'evaluation_config': {'judge_model': 'primary'}}
            (root / 'base.json').write_text(json.dumps(base))
            retry = {'engine': 'ragas', 'scores': [{'question': 'q', 'faithfulness': 0.0}]}
            (root / 'retry.json').write_text(json.dumps(retry))
            argv = ['merge', '--base', str(root / 'base.json'), '--retry', str(root / 'retry.json'),
                    '--retry-judge', 'retry', '--out-json', str(root / 'out.json'),
                    '--out-csv', str(root / 'nested/out.csv')]
            with patch('sys.argv', argv), redirect_stdout(io.StringIO()):
                main()
            report = json.loads((root / 'out.json').read_text())
            self.assertEqual(report['average']['faithfulness'], 0.0)
            self.assertIsNone(report['average']['answer_relevancy'])
            self.assertIsNone(report['composite'])
            self.assertEqual(report['complete_count'], 0)
            self.assertEqual(report['metric_judge_counts']['faithfulness'], {'primary': 0, 'retry': 1})
            retry['engine'] = 'llm_judge_fallback'
            (root / 'retry.json').write_text(json.dumps(retry))
            with patch('sys.argv', argv), self.assertRaisesRegex(ValueError, 'Ragas engine'):
                main()

    def test_html_extraction_preserves_medical_text(self):
        from scripts.extract_msd import clean_msd_html
        text = clean_msd_html('<main>HIV virus [说明]\n5\n120/80</main>')
        for expected in ['HIV virus', '[说明]', '5', '120/80']:
            self.assertIn(expected, text)

    def test_evaluation_report_missing_scores_and_relative_output(self):
        from scripts.evaluate_rag import print_and_save
        import os
        with tempfile.TemporaryDirectory() as directory:
            old = os.getcwd()
            try:
                os.chdir(directory)
                print_and_save({'metrics': ['faithfulness'], 'average': {'faithfulness': None},
                                'scores': []}, 'report.json', 'report.csv')
                self.assertTrue(Path('report.json').exists())
            finally:
                os.chdir(old)


if __name__ == '__main__':
    unittest.main()
